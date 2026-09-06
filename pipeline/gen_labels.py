"""Regenera pipeline/labels.py a partir das planilhas de codificação do IBGE.

Lê apenas as planilhas de código/rótulo (públicas, sem microdados individuais):
Codificação_Município, Codificação_País, Municípios_BR_Recortes_Territoriais_2022,
e o layout de acesso controlado (só cabeçalhos/categorias, não os CSVs de pessoas).

Uso: python pipeline/gen_labels.py <pasta_microdados>
"""
import sys, re, pathlib
from pprint import pformat
import openpyxl


def main(raw_dir: str) -> None:
    raw = pathlib.Path(raw_dir)
    out: dict[str, object] = {}

    # --- municípios (código -> nome/UF/sigla)
    wb = openpyxl.load_workbook(raw / "Codificação_Município CD2022.xlsx", read_only=True)
    mun: dict[str, dict] = {}
    for i, row in enumerate(wb.worksheets[0].iter_rows(values_only=True)):
        if i < 2 or row[2] is None:
            continue
        mun[str(row[2]).zfill(7)] = {
            "nome": row[3],
            "uf": str(row[1]).zfill(2) if row[1] else None,
            "sigla": row[0],
        }
    out["MUNICIPIOS"] = mun

    # --- países
    wb = openpyxl.load_workbook(raw / "Codificação_País CD2022.xlsx", read_only=True)
    pais: dict[str, str] = {}
    for i, row in enumerate(wb.worksheets[0].iter_rows(values_only=True)):
        if i < 2 or row[0] is None:
            continue
        pais[str(row[0])] = row[1]
    out["PAISES"] = pais

    # --- recortes territoriais (RGI, RGInt, meso, micro, RM/RIDE, aglomeração, concentração urbana)
    wb = openpyxl.load_workbook(raw / "Municípios_BR_Recortes_Territoriais_2022.xlsx", read_only=True)
    rows = list(wb.worksheets[0].iter_rows(values_only=True))
    hdr = rows[0]
    rec: dict[str, dict] = {}
    for r in rows[1:]:
        d = dict(zip(hdr, r))
        cd = str(d["CD_MUN"]).zfill(7)
        clean = {}
        for k, v in d.items():
            if k is None or k == "CD_MUN":
                continue
            if isinstance(v, float) and v.is_integer() and "cod" in k.lower():
                v = str(int(v))
            clean[k] = v
        rec[cd] = clean
    out["RECORTES"] = rec

    # --- categorias das variáveis (código -> rótulo) a partir do layout de acesso controlado
    wb = openpyxl.load_workbook(raw / "Layout Microdados CD2022 - acesso Controlado.xlsx", read_only=True)
    cats: dict[str, dict] = {}
    pat = re.compile(r"(?:(?<=\s)|^)(\d{1,3})\s*[-–]\s*(.+?)(?=\s+\d{1,3}\s*[-–]\s|\s*$)")
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            if not row or not isinstance(row[0], str) or not re.match(r"^[PDMF]\d{4}$", row[0]):
                continue
            var, nome = row[0], str(row[1]).replace("\n", " ")
            pairs = pat.findall(nome)
            if len(pairs) >= 2:
                desc = nome.split(pairs[0][0] + " ")[0].strip()[:120]
                cats[var] = {
                    "descricao": desc,
                    "categorias": {c: l.strip().rstrip(".") for c, l in pairs},
                }
    out["CATEGORIAS"] = cats

    out["UF"] = {
        "11": "Rondônia", "12": "Acre", "13": "Amazonas", "14": "Roraima", "15": "Pará",
        "16": "Amapá", "17": "Tocantins", "21": "Maranhão", "22": "Piauí", "23": "Ceará",
        "24": "Rio Grande do Norte", "25": "Paraíba", "26": "Pernambuco", "27": "Alagoas",
        "28": "Sergipe", "29": "Bahia", "31": "Minas Gerais", "32": "Espírito Santo",
        "33": "Rio de Janeiro", "35": "São Paulo", "41": "Paraná", "42": "Santa Catarina",
        "43": "Rio Grande do Sul", "50": "Mato Grosso do Sul", "51": "Mato Grosso",
        "52": "Goiás", "53": "Distrito Federal",
    }

    header = (
        '"""Rótulos e tabelas de códigos do Censo 2022.\n\n'
        "Gerado por `pipeline/gen_labels.py` a partir das planilhas públicas de codificação\n"
        "do IBGE (código de município/país/recortes territoriais e categorias do layout).\n"
        "Não contém microdados individuais. Não editar à mão -- rode o gerador novamente.\n"
        '"""\n\n'
    )
    dest = pathlib.Path("pipeline/labels.py")
    with dest.open("w", encoding="utf-8") as f:
        f.write(header)
        for k, v in out.items():
            f.write(f"{k} = {pformat(v, width=100, sort_dicts=False)}\n\n")
    print(f"labels.py: {dest.stat().st_size // 1024} KB; "
          f"{len(mun)} municípios; {len(pais)} países; "
          f"{len(cats)} variáveis categóricas; {len(rec)} recortes")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print("uso: python pipeline/gen_labels.py <pasta_microdados>", file=sys.stderr)
        raise SystemExit(1)
    main(sys.argv[1])
